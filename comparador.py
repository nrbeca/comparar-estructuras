"""
Comparador de Estructura Programática 2027 vs 2026 — SADER (Ramo 08)

Llave de comparación (UR y partida siempre amarradas):
    UR | Finalidad | Función | Subfunción | Actividad Institucional | Modalidad+Pp | Partida

Una llave de 2027 que no aparece en 2026 se marca como NUEVA y se explica el motivo:
    1. UR nueva                       -> la UR no existe en 2026
    2. Pp nuevo para la UR            -> la UR no tenía ese Modalidad+Pp en 2026
    3. Estructura nueva para la UR    -> la UR tenía el Pp, pero con otra FI/FN/SF/AI
    4. Partida nueva para la UR       -> la estructura existía, pero la UR nunca usó esa partida
    5. Partida en otra estructura     -> la UR sí usó la partida en 2026, pero en otra estructura
"""
from __future__ import annotations

import io
import re
import unicodedata

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------------------
# Campos y sinónimos para detectar columnas automáticamente (MAP, SICOP, bases PEF, etc.)
# --------------------------------------------------------------------------------------
CAMPOS = ["UR", "FI", "FN", "SF", "AI", "MOD", "PP", "PARTIDA", "IMPORTE"]

ETIQUETAS = {
    "UR": "UR",
    "FI": "Finalidad",
    "FN": "Función",
    "SF": "Subfunción",
    "AI": "Actividad Institucional",
    "MOD": "Modalidad",
    "PP": "Programa Presupuestario",
    "PARTIDA": "Partida",
    "IMPORTE": "Importe",
}

SINONIMOS = {
    "UR": ["ur", "unidadresponsable", "idunidad", "unidad", "cveur", "claveur", "idur"],
    "FI": ["finalidad", "fi", "idfinalidad", "cvefinalidad"],
    "FN": ["funcion", "fn", "fun", "idfuncion", "cvefuncion"],
    "SF": ["subfuncion", "sf", "sfu", "idsubfuncion", "cvesubfuncion"],
    "AI": ["actividadinstitucional", "ai", "idai", "actinst", "actividad", "cveai"],
    "MOD": ["modalidad", "mod", "idmodalidad", "cvemodalidad", "modpp"],
    "PP": ["programapresupuestario", "pp", "programa", "idpp", "idprograma", "cvepp", "prog"],
    "PARTIDA": ["partida", "partidaespecifica", "idpartida", "cvepartida", "objetodelgasto", "og", "ppartida"],
    "IMPORTE": ["importe", "monto", "aprobado", "original", "modificado", "modificadoanual",
                "pef", "ppef", "asignado", "presupuesto", "montoaprobado"],
}

# Para construir la partida cuando la base viene desagregada (MAP / SICOP)
PARTIDA_PARTES = {
    "CAPITULO": ["capitulo", "cap"],
    "CONCEPTO": ["concepto"],
    "GENERICA": ["partidagenerica", "generica"],
    "ESPECIFICA": ["partidaespecifica", "especifica"],
}

ESTRUCTURA = ["UR", "FI", "FN", "SF", "AI", "MOD", "PP"]
LLAVE = ESTRUCTURA + ["PARTIDA"]

ROJO = "9F2241"
CAFE = "BC955C"
CREMA = "F7F2EA"


def _norm(txt) -> str:
    s = unicodedata.normalize("NFKD", str(txt)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", s)


# --------------------------------------------------------------------------------------
# Lectura
# --------------------------------------------------------------------------------------
def hojas_de(archivo) -> list[str]:
    nombre = getattr(archivo, "name", str(archivo)).lower()
    if nombre.endswith((".csv", ".txt")):
        return ["(CSV)"]
    _rebobinar(archivo)
    xl = pd.ExcelFile(archivo)
    return xl.sheet_names


def _rebobinar(archivo):
    if hasattr(archivo, "seek"):
        archivo.seek(0)


def _puntaje_encabezado(fila) -> int:
    vals = {_norm(v) for v in fila if pd.notna(v)}
    todos = {s for lst in SINONIMOS.values() for s in lst} | {s for lst in PARTIDA_PARTES.values() for s in lst}
    return len(vals & todos)


def leer_base(archivo, hoja: str | None = None) -> pd.DataFrame:
    """Lee CSV o Excel detectando la fila de encabezados (la de más coincidencias)."""
    nombre = getattr(archivo, "name", str(archivo)).lower()
    _rebobinar(archivo)
    if nombre.endswith((".csv", ".txt")):
        crudo = None
        for enc in ("utf-8-sig", "latin-1"):
            try:
                _rebobinar(archivo)
                crudo = pd.read_csv(archivo, header=None, dtype=str, encoding=enc,
                                    sep=None, engine="python")
                break
            except UnicodeDecodeError:
                continue
    else:
        crudo = pd.read_excel(archivo, sheet_name=hoja or 0, header=None, dtype=object)

    fila_enc = max(range(min(20, len(crudo))), key=lambda i: _puntaje_encabezado(crudo.iloc[i]))
    enc = [str(c).strip() if pd.notna(c) else f"col_{i}" for i, c in enumerate(crudo.iloc[fila_enc])]
    # nombres duplicados (p.ej. 'UR' y 'Partida' repetidas al final de la base 2027)
    vistos, cols = {}, []
    for c in enc:
        if c in vistos:
            vistos[c] += 1
            cols.append(f"{c}.{vistos[c]}")
        else:
            vistos[c] = 0
            cols.append(c)
    df = crudo.iloc[fila_enc + 1:].copy()
    df.columns = cols
    df = df.dropna(how="all")
    # quita filas de totales tipo tabla dinámica
    primera = df.columns[0]
    df = df[~df[primera].astype(str).str.contains("total", case=False, na=False)]
    return df.reset_index(drop=True)


def detectar_columnas(df: pd.DataFrame) -> dict:
    """Devuelve {campo: nombre_columna} con la primera coincidencia por sinónimo."""
    normalizadas = {c: _norm(re.sub(r"\.\d+$", "", c)) for c in df.columns}
    mapeo = {}
    for campo in CAMPOS:
        for sin in SINONIMOS[campo]:
            hit = next((c for c, n in normalizadas.items() if n == sin and c not in mapeo.values()), None)
            if hit:
                mapeo[campo] = hit
                break
    # Partida desagregada (CAPITULO+CONCEPTO+GENERICA+ESPECIFICA)
    if "PARTIDA" not in mapeo or _norm(mapeo.get("PARTIDA", "")) == "partidaespecifica":
        partes = {}
        for p, sins in PARTIDA_PARTES.items():
            hit = next((c for c, n in normalizadas.items() if n in sins), None)
            if hit:
                partes[p] = hit
        if len(partes) == 4:
            mapeo.pop("PARTIDA", None)
            mapeo["_PARTES"] = partes
    return mapeo


# --------------------------------------------------------------------------------------
# Normalización
# --------------------------------------------------------------------------------------
def _limpio(v) -> str:
    if pd.isna(v):
        return ""
    s = str(v).strip()
    if re.fullmatch(r"-?\d+\.0+", s):
        s = s.split(".")[0]
    return s


def _pad(v, n) -> str:
    s = _limpio(v)
    return s.zfill(n) if s.isdigit() else s.upper()


def _sin_cero(v) -> str:
    s = _limpio(v)
    return str(int(s)) if s.isdigit() else s.upper()


def normalizar(df: pd.DataFrame, mapeo: dict) -> pd.DataFrame:
    out = pd.DataFrame(index=df.index)
    out["UR"] = df[mapeo["UR"]].map(lambda v: _pad(v, 3))
    out["FI"] = df[mapeo["FI"]].map(_sin_cero)
    out["FN"] = df[mapeo["FN"]].map(_sin_cero)
    out["SF"] = df[mapeo["SF"]].map(lambda v: _pad(v, 2))
    out["AI"] = df[mapeo["AI"]].map(lambda v: _pad(v, 3))
    out["MOD"] = df[mapeo["MOD"]].map(lambda v: _limpio(v).upper()) if mapeo.get("MOD") else ""
    out["PP"] = df[mapeo["PP"]].map(lambda v: _pad(v, 3))
    # Si la columna de Pp ya trae la modalidad (p.ej. 'S263'), separarla
    con_mod = out["PP"].str.match(r"^[A-Z]\d{3}$")
    out.loc[con_mod, "MOD"] = out.loc[con_mod, "PP"].str[0]
    out.loc[con_mod, "PP"] = out.loc[con_mod, "PP"].str[1:]

    if "_PARTES" in mapeo:
        p = mapeo["_PARTES"]
        # Acepta dígitos sueltos (1,3,1,01) o códigos completos (1000,1300,1310,13101)
        cap = df[p["CAPITULO"]].map(_limpio).str[:1]
        con = df[p["CONCEPTO"]].map(_limpio).map(lambda x: x[1] if len(x) >= 2 else x[-1:])
        gen = df[p["GENERICA"]].map(_limpio).map(lambda x: x[2] if len(x) >= 3 else x[-1:])
        esp = df[p["ESPECIFICA"]].map(_limpio).map(lambda x: x[3:5] if len(x) >= 5 else x.zfill(2)[-2:])
        out["PARTIDA"] = cap + con + gen + esp
    else:
        out["PARTIDA"] = df[mapeo["PARTIDA"]].map(lambda v: _pad(v, 5))

    if mapeo.get("IMPORTE"):
        out["IMPORTE"] = pd.to_numeric(
            df[mapeo["IMPORTE"]].astype(str).str.replace(r"[,$\s]", "", regex=True), errors="coerce"
        ).fillna(0.0)
    else:
        out["IMPORTE"] = 0.0

    out = out[(out["UR"] != "") & (out["PARTIDA"] != "")]
    out["LLAVE_ESTRUCTURA"] = out[ESTRUCTURA[:-2]].agg("-".join, axis=1) + "-" + out["MOD"] + out["PP"]
    out["LLAVE"] = out["LLAVE_ESTRUCTURA"] + "-" + out["PARTIDA"]
    out["UR_PARTIDA"] = out["UR"] + "-" + out["PARTIDA"]
    out["UR_PP"] = out["UR"] + "-" + out["MOD"] + out["PP"]
    return out


# --------------------------------------------------------------------------------------
# Comparación
# --------------------------------------------------------------------------------------
def comparar(n27: pd.DataFrame, n26: pd.DataFrame) -> dict:
    s26 = {
        "UR": set(n26["UR"]),
        "UR_PP": set(n26["UR_PP"]),
        "EST": set(n26["UR"] + "|" + n26["LLAVE_ESTRUCTURA"]),
        "UR_PARTIDA": set(n26["UR_PARTIDA"]),
        "LLAVE": set(n26["LLAVE"]),
    }

    def motivo(r):
        if r["LLAVE"] in s26["LLAVE"]:
            return ""
        if r["UR"] not in s26["UR"]:
            return "1. UR nueva"
        if r["UR_PP"] not in s26["UR_PP"]:
            return "2. Pp nuevo para la UR"
        if r["UR"] + "|" + r["LLAVE_ESTRUCTURA"] not in s26["EST"]:
            return "3. Estructura nueva para la UR"
        if r["UR_PARTIDA"] not in s26["UR_PARTIDA"]:
            return "4. Partida nueva para la UR"
        return "5. Partida usada en 2026 pero en otra estructura"

    cols = LLAVE + ["LLAVE_ESTRUCTURA", "LLAVE", "UR_PARTIDA", "UR_PP"]
    g27 = n27.groupby(cols, as_index=False)["IMPORTE"].sum().rename(columns={"IMPORTE": "IMPORTE_2027"})
    g26 = n26.groupby(cols, as_index=False)["IMPORTE"].sum().rename(columns={"IMPORTE": "IMPORTE_2026"})

    comp = g27.merge(g26[["LLAVE", "IMPORTE_2026"]], on="LLAVE", how="left")
    comp["ESTATUS"] = comp["IMPORTE_2026"].isna().map({True: "NUEVA 2027", False: "Existía en 2026"})
    comp["IMPORTE_2026"] = comp["IMPORTE_2026"].fillna(0.0)
    comp["MOTIVO"] = comp.apply(motivo, axis=1)
    comp["UR_PARTIDA_2026"] = comp["UR_PARTIDA"].isin(s26["UR_PARTIDA"]).map({True: "Sí", False: "No"})
    comp = comp.sort_values(["UR", "MOD", "PP", "AI", "PARTIDA"]).reset_index(drop=True)

    solo26 = g26[~g26["LLAVE"].isin(set(g27["LLAVE"]))].sort_values(["UR", "MOD", "PP", "PARTIDA"])

    res = comp.groupby("UR").agg(
        LLAVES_2027=("LLAVE", "count"),
        NUEVAS=("ESTATUS", lambda s: (s == "NUEVA 2027").sum()),
        IMPORTE_2027=("IMPORTE_2027", "sum"),
        IMPORTE_NUEVAS=("IMPORTE_2027", lambda s: s[comp.loc[s.index, "ESTATUS"] == "NUEVA 2027"].sum()),
    )
    res = res.join(solo26.groupby("UR")["LLAVE"].count().rename("SOLO_2026"), how="outer").fillna(0)
    res = res.reset_index()
    for c in ["LLAVES_2027", "NUEVAS", "SOLO_2026"]:
        res[c] = res[c].astype(int)

    return {"comparativo": comp, "solo_2026": solo26.reset_index(drop=True), "resumen": res,
            "n27": n27, "n26": n26}


# --------------------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------------------
_fino = Side(style="thin", color="D9D9D9")
_borde = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)
_fmt_dinero = '#,##0;[Red]-#,##0'


def _encabezado(ws, fila, titulos, anchos):
    for j, (t, w) in enumerate(zip(titulos, anchos), start=1):
        c = ws.cell(row=fila, column=j, value=t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=ROJO)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _borde
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[fila].height = 32


def _titulo(ws, texto, sub=""):
    ws["A1"] = texto
    ws["A1"].font = Font(bold=True, size=14, color=ROJO)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="7F7F7F")


def _escribe(ws, df, fila0, cols_dinero=()):
    for i, fila in enumerate(df.itertuples(index=False), start=fila0):
        for j, v in enumerate(fila, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = _borde
            if j in cols_dinero:
                c.number_format = _fmt_dinero


def generar_excel(r: dict, etiqueta_27="2027", etiqueta_26="2026") -> bytes:
    comp, solo26, res, n26 = r["comparativo"], r["solo_2026"], r["resumen"], r["n26"]
    wb = Workbook()

    # ---------------- Base 2026 (llaves) — referencia para las fórmulas ----------------
    ws26 = wb.active
    ws26.title = f"Base {etiqueta_26}"
    _titulo(ws26, f"Base {etiqueta_26} normalizada (referencia de las fórmulas)")
    g26 = n26.groupby(LLAVE + ["LLAVE"], as_index=False)["IMPORTE"].sum()
    t26 = ["LLAVE", "UR", "Finalidad", "Función", "Subfunción", "AI", "Mod", "Pp", "Partida", "UR-Partida", f"Importe {etiqueta_26}"]
    _encabezado(ws26, 4, t26, [30, 7, 9, 9, 10, 7, 6, 6, 9, 13, 16])
    d26 = g26[["LLAVE"] + LLAVE + ["IMPORTE"]].copy()
    d26.insert(9, "UR_PARTIDA", d26["UR"] + "-" + d26["PARTIDA"])
    _escribe(ws26, d26, 5, cols_dinero=(11,))
    n_26 = len(d26)
    ult26 = 4 + n_26
    ws26.freeze_panes = "B5"
    ws26.auto_filter.ref = f"A4:K{ult26}"
    rng_llave26 = f"'Base {etiqueta_26}'!$A$5:$A${max(ult26, 5)}"
    rng_imp26 = f"'Base {etiqueta_26}'!$K$5:$K${max(ult26, 5)}"
    rng_urp26 = f"'Base {etiqueta_26}'!$J$5:$J${max(ult26, 5)}"

    # ---------------- Comparativo (hoja principal) ----------------
    ws = wb.create_sheet(f"Comparativo {etiqueta_27}", 0)
    _titulo(ws, f"Estructura programática {etiqueta_27} vs {etiqueta_26} — UR y partida amarradas",
            "Llave = UR-FI-FN-SF-AI-Mod+Pp-Partida.  ESTATUS e Importe 2026 son fórmulas contra la hoja "
            f"'Base {etiqueta_26}'.")
    titulos = ["LLAVE", "UR", "Finalidad", "Función", "Subfunción", "AI", "Mod", "Pp", "Partida",
               f"Importe {etiqueta_27}", f"Importe {etiqueta_26}", "Diferencia", "ESTATUS",
               "Motivo (si es nueva)", f"¿UR-Partida existía en {etiqueta_26}?"]
    _encabezado(ws, 4, titulos, [30, 7, 9, 9, 10, 7, 6, 6, 9, 16, 16, 16, 16, 40, 16])
    for i, row in enumerate(comp.itertuples(index=False), start=5):
        vals = [row.LLAVE, row.UR, row.FI, row.FN, row.SF, row.AI, row.MOD, row.PP, row.PARTIDA,
                float(row.IMPORTE_2027),
                f"=SUMIF({rng_llave26},A{i},{rng_imp26})",
                f"=J{i}-K{i}",
                f'=IF(COUNTIF({rng_llave26},A{i})>0,"Existía en {etiqueta_26}","NUEVA {etiqueta_27}")',
                row.MOTIVO,
                f'=IF(COUNTIF({rng_urp26},B{i}&"-"&I{i})>0,"Sí","No")']
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = _borde
            if j in (10, 11, 12):
                c.number_format = _fmt_dinero
    ult = 4 + len(comp)
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:O{ult}"
    ws.conditional_formatting.add(
        f"A5:O{ult}",
        FormulaRule(formula=[f'LEFT($M5,5)="NUEVA"'], fill=PatternFill("solid", fgColor="F4D6DC"),
                    font=Font(color=ROJO, bold=True)))
    # Totales
    t = ult + 1
    ws.cell(row=t, column=9, value="TOTAL").font = Font(bold=True)
    for col in (10, 11, 12):
        L = get_column_letter(col)
        c = ws.cell(row=t, column=col, value=f"=SUBTOTAL(9,{L}5:{L}{ult})")
        c.font = Font(bold=True)
        c.number_format = _fmt_dinero
        c.fill = PatternFill("solid", fgColor=CREMA)

    # ---------------- Solo nuevas ----------------
    wn = wb.create_sheet(f"Nuevas {etiqueta_27}", 1)
    nuevas = comp[comp["ESTATUS"] == "NUEVA 2027"]
    _titulo(wn, f"Llaves {etiqueta_27} que nunca se usaron en {etiqueta_26}", f"{len(nuevas):,} llaves")
    tn = ["LLAVE", "UR", "Finalidad", "Función", "Subfunción", "AI", "Mod", "Pp", "Partida",
          f"Importe {etiqueta_27}", "Motivo", f"¿UR-Partida existía en {etiqueta_26}?"]
    _encabezado(wn, 4, tn, [30, 7, 9, 9, 10, 7, 6, 6, 9, 16, 40, 16])
    _escribe(wn, nuevas[["LLAVE"] + LLAVE + ["IMPORTE_2027", "MOTIVO", "UR_PARTIDA_2026"]], 5, cols_dinero=(10,))
    wn.freeze_panes = "B5"
    wn.auto_filter.ref = f"A4:L{4 + max(len(nuevas), 1)}"

    # ---------------- Solo 2026 ----------------
    wsol = wb.create_sheet(f"Solo {etiqueta_26}", 2)
    _titulo(wsol, f"Llaves que existían en {etiqueta_26} y no aparecen en {etiqueta_27}", f"{len(solo26):,} llaves")
    ts = ["LLAVE", "UR", "Finalidad", "Función", "Subfunción", "AI", "Mod", "Pp", "Partida", f"Importe {etiqueta_26}"]
    _encabezado(wsol, 4, ts, [30, 7, 9, 9, 10, 7, 6, 6, 9, 16])
    _escribe(wsol, solo26[["LLAVE"] + LLAVE + ["IMPORTE_2026"]], 5, cols_dinero=(10,))
    wsol.freeze_panes = "B5"
    wsol.auto_filter.ref = f"A4:J{4 + max(len(solo26), 1)}"

    # ---------------- Resumen por UR ----------------
    wr = wb.create_sheet("Resumen por UR", 0)
    _titulo(wr, f"Resumen por UR — estructura programática {etiqueta_27} vs {etiqueta_26}",
            "Conteos con fórmula sobre la hoja Comparativo; motivos al pie.")
    cname = f"'Comparativo {etiqueta_27}'"
    tr = ["UR", f"Llaves {etiqueta_27}", "Nuevas", "% nuevas", f"Importe {etiqueta_27}",
          "Importe en llaves nuevas", f"Llaves solo {etiqueta_26}"]
    _encabezado(wr, 4, tr, [8, 13, 10, 10, 18, 20, 14])
    for i, ur in enumerate(res["UR"], start=5):
        vals = [ur,
                f"=COUNTIF({cname}!$B$5:$B${ult},A{i})",
                f'=COUNTIFS({cname}!$B$5:$B${ult},A{i},{cname}!$M$5:$M${ult},"NUEVA*")',
                f"=IF(B{i}=0,0,C{i}/B{i})",
                f"=SUMIF({cname}!$B$5:$B${ult},A{i},{cname}!$J$5:$J${ult})",
                f'=SUMIFS({cname}!$J$5:$J${ult},{cname}!$B$5:$B${ult},A{i},{cname}!$M$5:$M${ult},"NUEVA*")',
                int(res.loc[res["UR"] == ur, "SOLO_2026"].iloc[0])]
        for j, v in enumerate(vals, start=1):
            c = wr.cell(row=i, column=j, value=v)
            c.border = _borde
            c.number_format = {4: "0.0%", 5: _fmt_dinero, 6: _fmt_dinero}.get(j, "General")
    ur_ult = 4 + len(res)
    t = ur_ult + 1
    wr.cell(row=t, column=1, value="TOTAL").font = Font(bold=True)
    for col in (2, 3, 5, 6, 7):
        L = get_column_letter(col)
        c = wr.cell(row=t, column=col, value=f"=SUM({L}5:{L}{ur_ult})")
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=CREMA)
        if col in (5, 6):
            c.number_format = _fmt_dinero
    c = wr.cell(row=t, column=4, value=f"=IF(B{t}=0,0,C{t}/B{t})")
    c.number_format = "0.0%"
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor=CREMA)
    wr.conditional_formatting.add(
        f"A5:G{ur_ult}", FormulaRule(formula=["$C5>0"], fill=PatternFill("solid", fgColor="F4D6DC")))
    wr.freeze_panes = "B5"

    # Motivos
    motivos = ["1. UR nueva", "2. Pp nuevo para la UR", "3. Estructura nueva para la UR",
               "4. Partida nueva para la UR", "5. Partida usada en 2026 pero en otra estructura"]
    for j, (txt, w) in enumerate([("Nuevas por motivo", 48), ("Llaves", 10)], start=9):
        c = wr.cell(row=4, column=j, value=txt)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=CAFE)
        c.border = _borde
        c.alignment = Alignment(horizontal="center", vertical="center")
        wr.column_dimensions[get_column_letter(j)].width = w
    for k, m in enumerate(motivos, start=5):
        wr.cell(row=k, column=9, value=m).border = _borde
        c = wr.cell(row=k, column=10, value=f'=COUNTIF({cname}!$N$5:$N${ult},I{k})')
        c.border = _borde
    c = wr.cell(row=10, column=9, value="TOTAL")
    c.font = Font(bold=True)
    c = wr.cell(row=10, column=10, value="=SUM(J5:J9)")
    c.font = Font(bold=True)
    for ws_ in wb.worksheets:
        ws_.sheet_view.showGridLines = False

    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
